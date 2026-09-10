import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, colorchooser, font
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os


# =========================================================
# APP
# =========================================================

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

app = ctk.CTk()

app.title("Untitled - Zoffice Documents")
app.geometry("1100x700")
app.minsize(800, 500)

current_file = None


# =========================================================
# VARIABLES
# =========================================================

font_family = tk.StringVar(value="Arial")
font_size = tk.StringVar(value="14")


# =========================================================
# FONT HELPERS
# =========================================================

def get_font_from_position(index):

    tags = editor.tag_names(index)

    family = "Arial"
    size = 14
    bold = False
    italic = False
    underline = False
    color = "#000000"

    for tag in tags:

        if tag.startswith("font_"):
            family = tag.replace(
                "font_",
                ""
            ).replace(
                "_",
                " "
            )

        elif tag.startswith("size_"):

            try:
                size = int(
                    tag.replace(
                        "size_",
                        ""
                    )
                )
            except:
                pass

        elif tag == "bold":
            bold = True

        elif tag == "italic":
            italic = True

        elif tag == "underline":
            underline = True

        elif tag.startswith("color_"):

            color = "#" + tag.replace(
                "color_",
                ""
            )

    return {
        "family": family,
        "size": size,
        "bold": bold,
        "italic": italic,
        "underline": underline,
        "color": color
    }


def apply_font_tag(start, end, family):

    tag = "font_" + family.replace(
        " ",
        "_"
    )

    editor.tag_configure(
        tag,
        font=font.Font(
            family=family,
            size=int(font_size.get())
        )
    )

    editor.tag_add(
        tag,
        start,
        end
    )


def apply_size_tag(start, end, size):

    tag = "size_" + str(size)

    editor.tag_configure(
        tag,
        font=font.Font(
            family=font_family.get(),
            size=int(size)
        )
    )

    editor.tag_add(
        tag,
        start,
        end
    )


# =========================================================
# TITLE
# =========================================================

def update_title():

    if current_file:

        app.title(
            f"{os.path.basename(current_file)} "
            f"- Zoffice Documents"
        )

    else:

        app.title(
            "Untitled - Zoffice Documents"
        )


# =========================================================
# NEW FILE
# =========================================================

def new_file():

    global current_file

    editor.delete(
        "1.0",
        "end"
    )

    current_file = None

    update_title()
    update_status()

    editor.focus()


# =========================================================
# CLEAR FORMATTING
# =========================================================

def clear_formatting():

    for tag in editor.tag_names():

        if tag != "sel":

            try:
                editor.tag_remove(
                    tag,
                    "1.0",
                    "end"
                )
            except:
                pass


# =========================================================
# OPEN TXT
# =========================================================

def open_txt(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8"
    ) as file:

        text = file.read()

    editor.delete(
        "1.0",
        "end"
    )

    clear_formatting()

    editor.insert(
        "1.0",
        text
    )

    update_status()


# =========================================================
# OPEN DOCX
# =========================================================

def open_docx(file_path):

    document = Document(file_path)

    editor.delete(
        "1.0",
        "end"
    )

    clear_formatting()

    for paragraph in document.paragraphs:

        paragraph_start = editor.index(
            "end-1c"
        )

        # -----------------------------------------
        # Insert Word runs
        # -----------------------------------------

        for run in paragraph.runs:

            if not run.text:
                continue

            start = editor.index(
                "end-1c"
            )

            editor.insert(
                "end",
                run.text
            )

            end = editor.index(
                "end-1c"
            )

            # -----------------------------------------
            # BOLD
            # -----------------------------------------

            if run.bold:

                editor.tag_add(
                    "bold",
                    start,
                    end
                )

            # -----------------------------------------
            # ITALIC
            # -----------------------------------------

            if run.italic:

                editor.tag_add(
                    "italic",
                    start,
                    end
                )

            # -----------------------------------------
            # UNDERLINE
            # -----------------------------------------

            if run.underline:

                editor.tag_add(
                    "underline",
                    start,
                    end
                )

            # -----------------------------------------
            # FONT FAMILY
            # -----------------------------------------

            if run.font.name:

                family = run.font.name

                tag = (
                    "font_"
                    + family.replace(
                        " ",
                        "_"
                    )
                )

                editor.tag_configure(
                    tag,
                    font=font.Font(
                        family=family,
                        size=14
                    )
                )

                editor.tag_add(
                    tag,
                    start,
                    end
                )

            # -----------------------------------------
            # FONT SIZE
            # -----------------------------------------

            if run.font.size:

                size = int(
                    run.font.size.pt
                )

                tag = "size_" + str(size)

                editor.tag_configure(
                    tag,
                    font=font.Font(
                        family="Arial",
                        size=size
                    )
                )

                editor.tag_add(
                    tag,
                    start,
                    end
                )

            # -----------------------------------------
            # TEXT COLOR
            # -----------------------------------------

            if (
                run.font.color
                and run.font.color.rgb
            ):

                rgb = str(
                    run.font.color.rgb
                )

                tag = "color_" + rgb

                editor.tag_configure(
                    tag,
                    foreground="#" + rgb
                )

                editor.tag_add(
                    tag,
                    start,
                    end
                )

        # -----------------------------------------
        # Paragraph alignment
        # -----------------------------------------

        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:

            editor.tag_add(
                "align_center",
                paragraph_start,
                f"{paragraph_start} lineend"
            )

        elif paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT:

            editor.tag_add(
                "align_right",
                paragraph_start,
                f"{paragraph_start} lineend"
            )

        else:

            editor.tag_add(
                "align_left",
                paragraph_start,
                f"{paragraph_start} lineend"
            )

        # -----------------------------------------
        # New paragraph
        # -----------------------------------------

        editor.insert(
            "end",
            "\n"
        )

    update_status()


# =========================================================
# OPEN FILE
# =========================================================

def open_file():

    global current_file

    file_path = filedialog.askopenfilename(

        title="Open Document",

        filetypes=[
            (
                "Word Document",
                "*.docx"
            ),
            (
                "Text File",
                "*.txt"
            ),
            (
                "All Files",
                "*.*"
            )
        ]
    )

    if not file_path:
        return

    try:

        extension = os.path.splitext(
            file_path
        )[1].lower()

        if extension == ".docx":

            open_docx(
                file_path
            )

        elif extension == ".txt":

            open_txt(
                file_path
            )

        else:

            # Try reading as text
            open_txt(
                file_path
            )

        current_file = file_path

        update_title()

        editor.focus()

    except Exception as error:

        from CTkMessagebox import CTkMessagebox

        CTkMessagebox(
            title="Zoffice",
            message=(
                "Failed to open file:\n\n"
                f"{error}"
            ),
            icon="cancel"
        )


# =========================================================
# APPLY RUN STYLE
# =========================================================

def apply_run_style(run, style):

    if not style:
        return

    family = style[0]
    size = style[1]
    bold = style[2]
    italic = style[3]
    underline = style[4]
    color = style[5]

    run.font.name = family

    run.font.size = Pt(
        size
    )

    run.bold = bold
    run.italic = italic
    run.underline = underline

    if color:

        color = color.replace(
            "#",
            ""
        )

        if len(color) == 6:

            run.font.color.rgb = (
                RGBColor.from_string(
                    color.upper()
                )
            )


# =========================================================
# SAVE DOCX
# =========================================================

def save_docx(file_path):

    document = Document()

    # Remove default paragraph

    first = document.paragraphs[0]

    first._element.getparent().remove(
        first._element
    )

    total_lines = int(
        editor.index(
            "end-1c"
        ).split(".")[0]
    )

    for line_number in range(
        1,
        total_lines + 1
    ):

        paragraph = document.add_paragraph()

        line_start = (
            f"{line_number}.0"
        )

        line_end = (
            f"{line_number}.end"
        )

        # -----------------------------------------
        # Alignment
        # -----------------------------------------

        tags = editor.tag_names(
            line_start
        )

        if "align_center" in tags:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

        elif "align_right" in tags:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
            )

        else:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
            )

        # -----------------------------------------
        # Text
        # -----------------------------------------

        text = editor.get(
            line_start,
            line_end
        )

        if not text:
            continue

        current_style = None
        current_text = ""

        # -----------------------------------------
        # Create runs
        # -----------------------------------------

        for column in range(
            len(text)
        ):

            index = (
                f"{line_number}."
                f"{column}"
            )

            style = get_font_from_position(
                index
            )

            style_tuple = (
                style["family"],
                style["size"],
                style["bold"],
                style["italic"],
                style["underline"],
                style["color"]
            )

            if (
                current_style is not None
                and style_tuple != current_style
            ):

                run = paragraph.add_run(
                    current_text
                )

                apply_run_style(
                    run,
                    current_style
                )

                current_text = ""

            current_text += text[column]

            current_style = style_tuple

        # -----------------------------------------
        # Last run
        # -----------------------------------------

        if current_text:

            run = paragraph.add_run(
                current_text
            )

            apply_run_style(
                run,
                current_style
            )

    document.save(
        file_path
    )


# =========================================================
# SAVE FILE
# =========================================================

def save_file():

    global current_file

    if current_file is None:

        save_as()

        return

    try:

        extension = os.path.splitext(
            current_file
        )[1].lower()

        if extension == ".docx":

            save_docx(
                current_file
            )

        else:

            content = editor.get(
                "1.0",
                "end-1c"
            )

            with open(
                current_file,
                "w",
                encoding="utf-8"
            ) as file:

                file.write(
                    content
                )

        update_title()

    except Exception as error:

        from CTkMessagebox import CTkMessagebox

        CTkMessagebox(
            title="Zoffice",
            message=(
                "Failed to save file:\n\n"
                f"{error}"
            ),
            icon="cancel"
        )


# =========================================================
# SAVE AS
# =========================================================

def save_as():

    global current_file

    file_path = filedialog.asksaveasfilename(

        title="Save Document",

        defaultextension=".docx",

        filetypes=[
            (
                "Word Document",
                "*.docx"
            ),
            (
                "Text File",
                "*.txt"
            )
        ]
    )

    if not file_path:
        return

    try:

        extension = os.path.splitext(
            file_path
        )[1].lower()

        if extension == ".txt":

            content = editor.get(
                "1.0",
                "end-1c"
            )

            with open(
                file_path,
                "w",
                encoding="utf-8"
            ) as file:

                file.write(
                    content
                )

        else:

            save_docx(
                file_path
            )

        current_file = file_path

        update_title()

    except Exception as error:

        from CTkMessagebox import CTkMessagebox

        CTkMessagebox(
            title="Zoffice",
            message=(
                "Failed to save file:\n\n"
                f"{error}"
            ),
            icon="cancel"
        )


# =========================================================
# BOLD
# =========================================================

def toggle_bold():

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        if "bold" in editor.tag_names(
            start
        ):

            editor.tag_remove(
                "bold",
                start,
                end
            )

        else:

            editor.tag_add(
                "bold",
                start,
                end
            )

    except tk.TclError:

        pass


# =========================================================
# ITALIC
# =========================================================

def toggle_italic():

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        if "italic" in editor.tag_names(
            start
        ):

            editor.tag_remove(
                "italic",
                start,
                end
            )

        else:

            editor.tag_add(
                "italic",
                start,
                end
            )

    except tk.TclError:

        pass


# =========================================================
# UNDERLINE
# =========================================================

def toggle_underline():

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        if "underline" in editor.tag_names(
            start
        ):

            editor.tag_remove(
                "underline",
                start,
                end
            )

        else:

            editor.tag_add(
                "underline",
                start,
                end
            )

    except tk.TclError:

        pass


# =========================================================
# CHANGE FONT
# =========================================================

def change_font(event=None):

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        apply_font_tag(
            start,
            end,
            font_family.get()
        )

    except tk.TclError:

        pass


# =========================================================
# CHANGE SIZE
# =========================================================

def change_size(event=None):

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        apply_size_tag(
            start,
            end,
            int(font_size.get())
        )

    except (tk.TclError, ValueError):

        pass


# =========================================================
# TEXT COLOR
# =========================================================

def change_color():

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        color = colorchooser.askcolor(
            title="Choose Text Color"
        )[1]

        if not color:
            return

        tag = (
            "color_"
            + color.replace(
                "#",
                ""
            )
        )

        editor.tag_configure(
            tag,
            foreground=color
        )

        editor.tag_add(
            tag,
            start,
            end
        )

    except tk.TclError:

        pass


# =========================================================
# SELECTED LINES
# =========================================================

def get_selected_lines():

    try:

        start = editor.index(
            "sel.first"
        )

        end = editor.index(
            "sel.last"
        )

        first_line = int(
            start.split(".")[0]
        )

        last_line = int(
            end.split(".")[0]
        )

        return (
            first_line,
            last_line
        )

    except tk.TclError:

        current = int(
            editor.index(
                "insert"
            ).split(".")[0]
        )

        return (
            current,
            current
        )


# =========================================================
# ALIGN LEFT
# =========================================================

def align_left():

    first, last = get_selected_lines()

    for line in range(
        first,
        last + 1
    ):

        editor.tag_add(
            "align_left",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_center",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_right",
            f"{line}.0",
            f"{line}.end"
        )


# =========================================================
# ALIGN CENTER
# =========================================================

def align_center():

    first, last = get_selected_lines()

    for line in range(
        first,
        last + 1
    ):

        editor.tag_add(
            "align_center",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_left",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_right",
            f"{line}.0",
            f"{line}.end"
        )


# =========================================================
# ALIGN RIGHT
# =========================================================

def align_right():

    first, last = get_selected_lines()

    for line in range(
        first,
        last + 1
    ):

        editor.tag_add(
            "align_right",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_left",
            f"{line}.0",
            f"{line}.end"
        )

        editor.tag_remove(
            "align_center",
            f"{line}.0",
            f"{line}.end"
        )


# =========================================================
# EDIT
# =========================================================

def cut_text():

    editor.event_generate(
        "<<Cut>>"
    )


def copy_text():

    editor.event_generate(
        "<<Copy>>"
    )


def paste_text():

    editor.event_generate(
        "<<Paste>>"
    )


def select_all(event=None):

    editor.tag_add(
        "sel",
        "1.0",
        "end"
    )

    return "break"


# =========================================================
# MENU FRAME
# =========================================================

menu_frame = ctk.CTkFrame(
    app,
    height=40,
    corner_radius=0
)

menu_frame.pack(
    fill="x",
    side="top"
)


# =========================================================
# FILE MENU BUTTON
# =========================================================

file_button = ctk.CTkButton(
    menu_frame,
    text="File",
    width=60,
    command=lambda:
        file_menu.tk_popup(
            file_button.winfo_rootx(),
            file_button.winfo_rooty()
            + file_button.winfo_height()
        )
)

file_button.pack(
    side="left",
    padx=5,
    pady=5
)


# =========================================================
# EDIT MENU BUTTON
# =========================================================

edit_button = ctk.CTkButton(
    menu_frame,
    text="Edit",
    width=60,
    command=lambda:
        edit_menu.tk_popup(
            edit_button.winfo_rootx(),
            edit_button.winfo_rooty()
            + edit_button.winfo_height()
        )
)

edit_button.pack(
    side="left",
    padx=5,
    pady=5
)


# =========================================================
# FILE MENU
# =========================================================

file_menu = tk.Menu(
    app,
    tearoff=0
)

file_menu.add_command(
    label="New",
    command=new_file
)

file_menu.add_command(
    label="Open",
    command=open_file
)

file_menu.add_command(
    label="Save",
    command=save_file
)

file_menu.add_command(
    label="Save As",
    command=save_as
)

file_menu.add_separator()

file_menu.add_command(
    label="Exit",
    command=app.destroy
)


# =========================================================
# EDIT MENU
# =========================================================

edit_menu = tk.Menu(
    app,
    tearoff=0
)

edit_menu.add_command(
    label="Cut",
    command=cut_text
)

edit_menu.add_command(
    label="Copy",
    command=copy_text
)

edit_menu.add_command(
    label="Paste",
    command=paste_text
)

edit_menu.add_separator()

edit_menu.add_command(
    label="Select All",
    command=select_all
)


# =========================================================
# TOOLBAR
# =========================================================

toolbar = ctk.CTkFrame(
    app,
    height=55,
    corner_radius=0
)

toolbar.pack(
    fill="x",
    side="top"
)


# =========================================================
# FONT
# =========================================================

font_box = ctk.CTkComboBox(
    toolbar,
    variable=font_family,
    values=[
        "Arial",
        "Calibri",
        "Times New Roman",
        "Courier New",
        "Verdana",
        "Georgia"
    ],
    width=150,
    command=change_font
)

font_box.pack(
    side="left",
    padx=(10, 5),
    pady=10
)


# =========================================================
# SIZE
# =========================================================

size_box = ctk.CTkComboBox(
    toolbar,
    variable=font_size,
    values=[
        "8",
        "10",
        "12",
        "14",
        "16",
        "18",
        "20",
        "24",
        "28",
        "32",
        "36",
        "48",
        "72"
    ],
    width=70,
    command=change_size
)

size_box.pack(
    side="left",
    padx=5,
    pady=10
)


# =========================================================
# BOLD BUTTON
# =========================================================

bold_button = ctk.CTkButton(
    toolbar,
    text="B",
    width=40,
    command=toggle_bold
)

bold_button.pack(
    side="left",
    padx=3
)


# =========================================================
# ITALIC BUTTON
# =========================================================

italic_button = ctk.CTkButton(
    toolbar,
    text="I",
    width=40,
    command=toggle_italic
)

italic_button.pack(
    side="left",
    padx=3
)


# =========================================================
# UNDERLINE BUTTON
# =========================================================

underline_button = ctk.CTkButton(
    toolbar,
    text="U",
    width=40,
    command=toggle_underline
)

underline_button.pack(
    side="left",
    padx=3
)


# =========================================================
# COLOR BUTTON
# =========================================================

color_button = ctk.CTkButton(
    toolbar,
    text="Text Color",
    width=90,
    command=change_color
)

color_button.pack(
    side="left",
    padx=5
)


# =========================================================
# SEPARATOR
# =========================================================

ctk.CTkLabel(
    toolbar,
    text="│",
    width=10
).pack(
    side="left"
)


# =========================================================
# ALIGNMENT BUTTONS
# =========================================================

ctk.CTkButton(
    toolbar,
    text="Left",
    width=55,
    command=align_left
).pack(
    side="left",
    padx=3
)


ctk.CTkButton(
    toolbar,
    text="Center",
    width=65,
    command=align_center
).pack(
    side="left",
    padx=3
)


ctk.CTkButton(
    toolbar,
    text="Right",
    width=60,
    command=align_right
).pack(
    side="left",
    padx=3
)


# =========================================================
# EDITOR FRAME
# =========================================================

editor_frame = ctk.CTkFrame(
    app,
    corner_radius=0
)

editor_frame.pack(
    fill="both",
    expand=True,
    padx=15,
    pady=(5, 10)
)


# =========================================================
# SCROLLBAR
# =========================================================

scrollbar = ctk.CTkScrollbar(
    editor_frame,
    orientation="vertical"
)

scrollbar.pack(
    side="right",
    fill="y"
)


# =========================================================
# TEXT EDITOR
# =========================================================

editor = tk.Text(
    editor_frame,

    wrap="word",

    undo=True,

    font=(
        "Arial",
        14
    ),

    bg="white",

    fg="black",

    insertbackground="black",

    selectbackground="#3399FF",

    padx=30,

    pady=25,

    relief="flat",

    borderwidth=0,

    yscrollcommand=scrollbar.set
)

editor.pack(
    fill="both",
    expand=True
)


scrollbar.configure(
    command=editor.yview
)


# =========================================================
# DEFAULT TAGS
# =========================================================

editor.tag_configure(
    "bold",
    font=(
        "Arial",
        14,
        "bold"
    )
)

editor.tag_configure(
    "italic",
    font=(
        "Arial",
        14,
        "italic"
    )
)

editor.tag_configure(
    "underline",
    underline=True
)


editor.tag_configure(
    "align_left",
    justify="left"
)

editor.tag_configure(
    "align_center",
    justify="center"
)

editor.tag_configure(
    "align_right",
    justify="right"
)


# =========================================================
# STATUS BAR
# =========================================================

status_bar = ctk.CTkLabel(
    app,
    text=(
        "Ln 1, Col 1"
        "     |     "
        "Zoffice Documents"
    ),
    anchor="w",
    height=25
)

status_bar.pack(
    fill="x",
    side="bottom",
    padx=10
)


# =========================================================
# UPDATE STATUS
# =========================================================

def update_status(event=None):

    cursor = editor.index(
        tk.INSERT
    )

    line, column = cursor.split(".")

    status_bar.configure(
        text=(
            f"Ln {line}, "
            f"Col {int(column) + 1}"
            f"     |     "
            f"Zoffice Documents"
        )
    )


editor.bind(
    "<KeyRelease>",
    update_status
)

editor.bind(
    "<ButtonRelease>",
    update_status
)


# =========================================================
# SHORTCUTS
# =========================================================

app.bind(
    "<Control-n>",
    lambda event: new_file()
)

app.bind(
    "<Control-o>",
    lambda event: open_file()
)

app.bind(
    "<Control-s>",
    lambda event: save_file()
)

app.bind(
    "<Control-Shift-S>",
    lambda event: save_as()
)

app.bind(
    "<Control-b>",
    lambda event: toggle_bold()
)

app.bind(
    "<Control-i>",
    lambda event: toggle_italic()
)

app.bind(
    "<Control-u>",
    lambda event: toggle_underline()
)

app.bind(
    "<Control-a>",
    select_all
)


# =========================================================
# START
# =========================================================

update_title()

editor.focus()

app.mainloop()